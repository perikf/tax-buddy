"""Workbook generation: CGT Summary, 1 STOCKS, Pool History, FX Detail, Data Quality."""
import io
import re
from datetime import date
from decimal import Decimal
from itertools import groupby

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.engine import SHORT_CLOSE
from backend.models import DisposalResult, FxRecord, HistoryGap, PoolEvent

GBP_FMT = "£#,##0.00"
DATE_FMT = "DD-MMM-YYYY"
RATE_FMT = "#,##0.000000"
QTY_FMT = "#,##0.########"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", start_color="312E81")
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)
AMBER_FILL = PatternFill("solid", start_color="FEF3C7")


def build_filename(period_label: str | None, start: date, end: date) -> str:
    if period_label and re.fullmatch(r"\d{4}/\d{2}", period_label):
        return f"TaxBuddy_CGT_{period_label.replace('/', '-')}.xlsx"
    return f"TaxBuddy_CGT_{start.isoformat()}_to_{end.isoformat()}.xlsx"


def _header_row(ws, row: int, headers: list[str]) -> None:
    for col, name in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=name)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _fmt(ws, row: int, formats: dict[int, str]) -> None:
    for col, fmt in formats.items():
        ws.cell(row=row, column=col).number_format = fmt


def _label_rows(ws, start_row: int, rows: list[tuple], bold_label: bool = False) -> None:
    for r, (label, value, fmt) in enumerate(rows, start_row):
        cell = ws.cell(row=r, column=1, value=label)
        if bold_label:
            cell.font = BOLD
        ws.cell(row=r, column=2, value=value)
        if fmt:
            _fmt(ws, r, {2: fmt})


def generate_excel(disposals: list[DisposalResult], pool_events: list[PoolEvent],
                   fx_records: list[FxRecord], gaps: list[HistoryGap],
                   period_label: str, start: date, end: date, warnings: list[str]) -> bytes:
    wb = Workbook()
    _summary_tab(wb.active, disposals, period_label, start, end, warnings, bool(gaps))
    if disposals:
        _stocks_tab(wb.create_sheet("1 STOCKS"), disposals)
    _pool_tab(wb.create_sheet("Pool History"), pool_events)
    _fx_tab(wb.create_sheet("FX Detail"), fx_records)
    if gaps:
        _gaps_tab(wb.create_sheet("Data Quality Warning"), gaps)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _summary_tab(ws, disposals, period_label, start, end, warnings, has_gaps) -> None:
    ws.title = "CGT Summary"
    ws["A1"] = "Tax Buddy — UK Capital Gains Tax Report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Period: {period_label}  ({start.isoformat()} to {end.isoformat()})"
    ws["A3"] = "Disposals matched per TCGA 1992: same-day, 30-day bed & breakfast, Section 104 holding."

    proceeds = sum((d.proceeds_gbp for d in disposals), Decimal(0))
    costs = sum((d.allowable_cost_gbp for d in disposals), Decimal(0))
    gains = sum((d.gain_loss_gbp for d in disposals if d.gain_loss_gbp > 0), Decimal(0))
    losses = sum((-d.gain_loss_gbp for d in disposals if d.gain_loss_gbp < 0), Decimal(0))
    n_disposals = len({(d.symbol, d.disposal_date) for d in disposals})

    _header_row(ws, 5, ["Product", "Disposals", "Proceeds", "Allowable Costs", "Gains", "Losses", "Net Gain/Loss"])
    for col, v in enumerate(["1 STOCKS", n_disposals, proceeds, costs, gains, losses, gains - losses], 1):
        ws.cell(row=6, column=col, value=v)
    for col, v in enumerate(["TOTAL", n_disposals, proceeds, costs, gains, losses, gains - losses], 1):
        cell = ws.cell(row=7, column=col, value=v)
        cell.font = BOLD
    for r in (6, 7):
        _fmt(ws, r, {c: GBP_FMT for c in range(3, 8)})

    ws["A9"] = "SA108 — Listed shares and securities"
    ws["A9"].font = BOLD
    _label_rows(ws, 10, [
        ("Box 23 — Number of disposals", n_disposals, None),
        ("Box 24 — Disposal proceeds", proceeds, GBP_FMT),
        ("Box 25 — Allowable costs (including purchase price)", costs, GBP_FMT),
        ("Box 26 — Gains in the year, before losses", gains, GBP_FMT),
        ("Box 27 — Losses in the year", losses, GBP_FMT),
    ])

    row = 16
    if has_gaps:
        ws.cell(row=row, column=1,
                value="⚠ Trade history looks incomplete — see the Data Quality Warning tab before filing.")
        ws.cell(row=row, column=1).fill = AMBER_FILL
        row += 1
    if warnings:
        ws.cell(row=row, column=1, value=f"{len(warnings)} processing warning(s):").font = BOLD
        for i, w in enumerate(warnings, row + 1):
            ws.cell(row=i, column=1, value=w)
    _widths(ws, [52, 16, 14, 16, 14, 14, 16])


def _stocks_tab(ws, disposals) -> None:
    total_gain = sum((d.gain_loss_gbp for d in disposals), Decimal(0))
    ws["A1"] = "1 STOCKS — Disposals"
    ws["A1"].font = TITLE_FONT
    _label_rows(ws, 2, [
        ("Disposal lots", len(disposals), None),
        ("Total proceeds", sum((d.proceeds_gbp for d in disposals), Decimal(0)), GBP_FMT),
        ("Total allowable costs", sum((d.allowable_cost_gbp for d in disposals), Decimal(0)), GBP_FMT),
        ("Net gain/loss", total_gain, GBP_FMT),
    ], bold_label=True)

    headers = ["Broker", "Product", "Domicile", "Ticker", "Description", "ISIN/Conid",
               "Trade Date", "Type", "Open/Close", "Currency", "Quantity", "Proceeds (CCY)",
               "FX Rate", "FX Rate Date", "Proceeds GBP", "Allowable Cost GBP",
               "Commission GBP", "Gain/Loss GBP", "Matching Rule", "Acquisition Date"]
    _header_row(ws, 7, headers)
    formats = {7: DATE_FMT, 11: QTY_FMT, 12: "#,##0.00", 13: RATE_FMT, 14: DATE_FMT,
               15: GBP_FMT, 16: GBP_FMT, 17: GBP_FMT, 18: GBP_FMT, 20: DATE_FMT}

    row = 8
    for (symbol, isin), group in groupby(disposals, key=lambda d: (d.symbol, d.isin_conid)):
        group = list(group)
        for d in group:
            values = ["IBKR", "Stocks", d.domicile, d.symbol, d.description, d.isin_conid,
                      d.disposal_date, "Buy (cover)" if d.matching_rule == SHORT_CLOSE else "Sell",
                      "Close", d.orig_currency, d.quantity_disposed, d.proceeds_orig,
                      d.fx_rate, d.fx_rate_date, d.proceeds_gbp, d.allowable_cost_gbp,
                      d.commission_gbp, d.gain_loss_gbp, d.matching_rule, d.acquisition_date]
            for col, v in enumerate(values, 1):
                ws.cell(row=row, column=col, value=v)
            _fmt(ws, row, formats)
            row += 1
        ws.cell(row=row, column=4, value=f"TOTAL {symbol}").font = BOLD
        total = ws.cell(row=row, column=18, value=sum((d.gain_loss_gbp for d in group), Decimal(0)))
        total.font, total.number_format = BOLD, GBP_FMT
        row += 2
    _widths(ws, [8, 8, 14, 10, 28, 14, 13, 11, 10, 9, 10, 12, 11, 13, 13, 15, 13, 13, 15, 14])


def _pool_tab(ws, events) -> None:
    _header_row(ws, 1, ["Date", "Asset", "Product", "Event", "Qty Δ", "Cost Δ GBP",
                        "Pool Qty", "Pool Cost", "Avg Cost/Unit"])
    for row, e in enumerate(events, 2):
        for col, v in enumerate([e.date, e.symbol, e.asset_category, e.event_type,
                                 e.quantity_delta, e.cost_delta_gbp, e.pool_quantity_after,
                                 e.pool_cost_after_gbp, e.average_cost_after_gbp], 1):
            ws.cell(row=row, column=col, value=v)
        _fmt(ws, row, {1: DATE_FMT, 5: QTY_FMT, 6: GBP_FMT, 7: QTY_FMT, 8: GBP_FMT, 9: GBP_FMT})
    _widths(ws, [13, 10, 9, 12, 11, 13, 11, 13, 13])


def _fx_tab(ws, records) -> None:
    _header_row(ws, 1, ["Trade Date", "Asset", "Type", "Currency", "Orig Amount",
                        "Rate", "Rate Date", "GBP Amount", "Note"])
    for row, r in enumerate(records, 2):
        for col, v in enumerate([r.trade_date, r.symbol, r.trade_type, r.currency,
                                 r.orig_amount, r.rate, r.rate_date, r.gbp_amount, r.note], 1):
            ws.cell(row=row, column=col, value=v)
        _fmt(ws, row, {1: DATE_FMT, 5: "#,##0.00", 6: RATE_FMT, 7: DATE_FMT, 8: GBP_FMT})
    _widths(ws, [13, 10, 7, 9, 13, 11, 13, 13, 30])


def _gaps_tab(ws, gaps) -> None:
    ws["A1"] = "⚠ Incomplete trade history detected — figures may be wrong. Upload statements covering ALL years."
    ws["A1"].font, ws["A1"].fill = BOLD, AMBER_FILL
    _header_row(ws, 3, ["Asset", "Product", "Pool Qty", "IBKR Held Qty", "Gap", "First-Trade-is-Sell"])
    for row, g in enumerate(gaps, 4):
        for col, v in enumerate([g.symbol, "Stocks", g.pool_qty,
                                 "n/a" if g.open_qty is None else g.open_qty,
                                 "n/a" if g.gap is None else g.gap,
                                 "YES" if g.first_trade_is_sell else ""], 1):
            ws.cell(row=row, column=col, value=v)
        _fmt(ws, row, {3: QTY_FMT, 4: QTY_FMT, 5: QTY_FMT})
    _widths(ws, [10, 9, 11, 14, 10, 18])
